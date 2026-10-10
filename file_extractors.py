"""Content-aware, bounded file readers for morogaro.
Uploads remain untrusted. Do not execute scripts, macros, PostScript or .blend
code. ZIP members are read in memory, never extractall(). Previews are actual
rasterizations/embedded images, never AI-generated replacement art.
"""
from __future__ import annotations
import base64, csv, hashlib, io, json, math, os, re, shutil, stat, struct, subprocess, tempfile, warnings, zipfile, zlib
from collections import Counter
from pathlib import Path, PurePosixPath
from html.parser import HTMLParser
from PIL import Image, ImageOps, ImageDraw, ImageFont, ImageStat
from defusedxml import ElementTree as ET
import fitz
from analysis_config import *
from analysis_engine import evidence, inspect_text, add_font, uniq, TOOLS, TYPES

TEXT_EXT = set('txt md markdown rst log css scss less html htm json jsonl yaml yml csv tsv xml svg js jsx ts tsx py java c cpp h hpp cs go rs rb php sql sh swift kt dart ipynb srt vtt'.split())
IMAGE_EXT = set('png jpg jpeg webp gif bmp tif tiff avif heic heif psd psb ico tga'.split())
VIDEO_EXT = set('mp4 mov webm mkv avi m4v mpg mpeg'.split())
AUDIO_EXT = set('mp3 wav m4a flac ogg aac opus'.split())
OFFICE_EXT = set('docx pptx xlsx xlsm odt ods odp hwpx'.split())
NEVER_AUTO = set('exe dll so dylib app dmg pkg bat com scr ttf otf woff woff2 pem key p12 pfx keystore env'.split())
KIND = {**TYPES, **{e:'영상' for e in VIDEO_EXT}, **{e:'오디오' for e in AUDIO_EXT}, **{e:'문서' for e in OFFICE_EXT}, **{e:'이미지' for e in IMAGE_EXT}, 'psd':'디자인 원본 파일','psb':'디자인 원본 파일','hwp':'문서','zip':'압축파일','sketch':'디자인 원본 파일','kra':'디자인 원본 파일','ora':'디자인 원본 파일','pages':'문서','key':'프레젠테이션','numbers':'데이터'}
TOOL_MAP = {**TOOLS,'psb':'Adobe Photoshop','hwp':'한글','hwpx':'한글','kra':'Krita'}

class LimitError(ValueError): pass


def sample_indices(total:int, limit:int) -> list[int]:
    """Cover beginning, middle AND end rather than silently taking only a cover."""
    if total <= limit:return list(range(total))
    if limit <= 1:return [0]
    return sorted(set(round(i*(total-1)/(limit-1)) for i in range(limit)))


def clean_name(name):
    return re.sub(r'[\x00-\x1f\x7f]', '', str(name)).replace('\\','/')[:400]


def decode_text(data:bytes) -> str:
    head=data[:256*1024]
    for codec in (['utf-16'] if head.startswith((b'\xff\xfe',b'\xfe\xff')) else ['utf-8-sig','cp949','utf-16']):
        try:return head.decode(codec)
        except UnicodeError:continue
    return head.decode('utf-8',errors='replace')


def bounded_text(text, limit=MAX_TEXT):
    """A strict budget including omission markers; even limit=0 is safe."""
    limit=max(0,int(limit))
    if len(text)<=limit:return text
    markers=('\n[중략: 중간 발췌]\n','\n[중략: 끝부분]\n')
    budget=limit-sum(map(len,markers))
    if budget<9:return text[:limit]
    first=budget//3;middle=budget//3;last=budget-first-middle
    start=max(first,len(text)//2-middle//2)
    return text[:first]+markers[0]+text[start:start+middle]+markers[1]+text[-last:]


def font_path():
    override=os.getenv('PREVIEW_FONT_PATH','')
    paths=[override,'/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc','/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf','/System/Library/Fonts/AppleSDGothicNeo.ttc','C:/Windows/Fonts/malgun.ttf']
    return next((p for p in paths if p and Path(p).is_file()),None)


def text_card(title, lines):
    """Fallback only for sheets. Clearly labelled data view, NOT an Excel screenshot."""
    image=Image.new('RGB',(1100,700),'white');d=ImageDraw.Draw(image)
    f=font_path();font=ImageFont.truetype(f,20) if f else ImageFont.load_default(size=20)
    heading=ImageFont.truetype(f,26) if f else ImageFont.load_default(size=26)
    d.text((30,24),title[:62],font=heading,fill='#262626')
    for i,line in enumerate(lines[:17]):
        y=78+i*33
        if i%2==0:d.rectangle((20,y-2,1080,y+30),fill='#F3F6F4')
        d.text((30,y),str(line)[:88],font=font,fill='#333333')
    d.text((30,666),'표 데이터 미리보기 · 원본의 시각 디자인/차트는 재현하지 않음',font=font,fill='#737373')
    return image


def safe_zip(data:bytes):
    z=zipfile.ZipFile(io.BytesIO(data));infos=z.infolist()
    if len(infos)>MAX_ZIP_ENTRIES or sum(i.file_size for i in infos)>MAX_EXPANDED_BYTES:
        z.close();raise LimitError('압축 해제 범위가 너무 커요.')
    if any(i.file_size>MAX_MEMBER or (i.file_size>1024*1024 and i.file_size/max(i.compress_size,1)>MAX_ZIP_RATIO) for i in infos):
        z.close();raise LimitError('압축 항목의 크기 또는 압축률 한도를 초과했어요.')
    return z


def is_safe_member(info):
    p=PurePosixPath(info.filename.replace('\\','/'))
    return not (p.is_absolute() or '..' in p.parts or ':' in info.filename or info.flag_bits&1 or stat.S_ISLNK(info.external_attr>>16) or len(info.filename)>400)


def public_capabilities():
    ffmpeg=shutil.which('ffmpeg');ffprobe=shutil.which('ffprobe')
    office=find_office()
    try:import pillow_heif;heif=True
    except ImportError:heif=False
    return dict(video=bool(ffmpeg and ffprobe),audio=bool(ffmpeg and ffprobe),officeRender=bool(office and office_enabled()),officeInstalled=bool(office),heic=heif,
                documents='PDF, DOCX, PPTX, XLSX, HWPX, ODT/ODS/ODP',archive='ZIP (최대 2단계)',
                limitations='보호·손상 파일, 비공개 바이너리 포맷, 외부 링크, 일부 코덱은 부분 확인 또는 미지원')


def find_office():
    explicit=os.getenv('SOFFICE_PATH','')
    if explicit and Path(explicit).is_file():return explicit
    return shutil.which('soffice') or shutil.which('libreoffice') or next((p for p in ['/Applications/LibreOffice.app/Contents/MacOS/soffice',r'C:\Program Files\LibreOffice\program\soffice.exe'] if Path(p).is_file()),None)


def scrub_env(home):
    # Never let a file converter inherit OPENAI_API_KEY or the user's other secrets.
    keep=('PATH','SYSTEMROOT','WINDIR','TMPDIR','TEMP','TMP','LANG','LC_ALL')
    env={k:v for k,v in os.environ.items() if k in keep}
    env.update(HOME=str(home),OMP_NUM_THREADS='1',OPENBLAS_NUM_THREADS='1')
    return env


class Extractor:
    def __init__(self,result,progress=None):
        self.result=result;self.progress=progress or (lambda *_:None);self.preview_bytes=0;self.expanded=0;self.conversions=0;self.audio_count=0
        self.previews_per_source=Counter();self.hashes={}

    def warn(self,s,message):
        if message not in s['warnings']:s['warnings'].append(message)

    def text(self,s,text,coverage):
        short=bounded_text(text);s['text']=short
        s['coverage']=coverage;s['status']='partial' if short!=text else 'read'
        if short!=text:self.warn(s,'긴 본문은 처음·중간·끝에서 합계 약 24,000자만 발췌했어요.')
        inspect_text(short,s,self.result)

    def capture(self,im,s,label,kind='image',page=None,member=None,seconds=None,sample_colors=True):
        if im.width*im.height>40_000_000:raise LimitError('이미지 해상도 한도를 초과했어요.')
        if min(im.size)<24:return
        rgba=ImageOps.exif_transpose(im).convert('RGBA')
        rgba.thumbnail((1120,1120))
        rgb=Image.new('RGB',rgba.size,'white');rgb.paste(rgba,mask=rgba.getchannel('A'))
        digest=hashlib.sha256(str(rgb.size).encode()+rgb.tobytes()).hexdigest()[:20]
        # Exact duplicate removal within a source; don't lose evidence from a different source.
        identity=(s['id'],digest)
        if identity in self.hashes:
            old=self.hashes[identity];old.setdefault('alsoAt',[])
            if label not in old['alsoAt'] and len(old['alsoAt'])<10:old['alsoAt'].append(label)
            return
        if len(self.result['thumbnails'])>=MAX_THUMBNAILS or self.previews_per_source[s['id']]>=12:
            self.warn(s,'썸네일 후보 한도에 도달해 나머지 캡처는 생략했어요.');s['status']='partial';return
        buf=io.BytesIO();rgb.save(buf,format='JPEG',quality=85,optimize=True);data=buf.getvalue()
        if self.preview_bytes+len(data)>MAX_PREVIEW_BYTES:
            self.warn(s,'썸네일 전체 용량 한도에 도달했어요.');s['status']='partial';return
        self.preview_bytes+=len(data);self.previews_per_source[s['id']]+=1
        # Real pixels determine simple automatic cover ranking; not semantic AI generation.
        gray=rgb.convert('L');gray.thumbnail((100,100));hist=gray.histogram();count=sum(hist)
        blank=sum(hist[:7])+sum(hist[248:]);texture=ImageStat.Stat(gray).stddev[0]
        score=round(min(1,texture/65)*.6+min(1,min(im.size)/600)*.3+(1-blank/max(count,1))*.1,3)
        thumb=dict(id=s['id']+'-'+digest,sourceId=s['id'],label=label,dataUrl='data:image/jpeg;base64,'+base64.b64encode(data).decode(),page=page,kind=kind,member=member,seconds=seconds,width=im.width,height=im.height,quality=score)
        self.hashes[identity]=thumb;self.result['thumbnails'].append(thumb)
        if sample_colors:
            sample=rgb.copy();sample.thumbnail((160,160));q=sample.quantize(colors=18,method=Image.Quantize.MEDIANCUT);pal=q.getpalette();out=[]
            for n,idx in sorted(q.getcolors() or [],reverse=True):
                c=tuple(pal[idx*3:idx*3+3])
                if n/(sample.width*sample.height)<.002:continue
                if any(sum((a-b)**2 for a,b in zip(c,p))<28**2 for p in out):continue
                out.append(c)
            self.result['colors']=uniq(self.result['colors']+['#%02X%02X%02X'%c for c in out])[:24]
            evidence(self.result,'colors',s,'캡처 픽셀의 대표 RGB 색상. 인쇄 잉크/ICC 값과 다를 수 있습니다.','measured',label)

    def pdf(self,data,s):
        with fitz.open(stream=data,filetype='pdf') as doc:
            if doc.needs_pass:raise ValueError('암호가 걸린 PDF는 내용을 읽지 않았어요.')
            pages=sample_indices(doc.page_count,MAX_TEXT_PAGES);thumbs=sample_indices(doc.page_count,MAX_PAGE_PREVIEWS);texts=[]
            s.update(pagesTotal=doc.page_count,pagesRead=len(pages),pageNumbersRead=[i+1 for i in pages],previewPages=[i+1 for i in thumbs])
            for i in sorted(set(pages)|set(thumbs)):
                page=doc[i]
                if i in pages:texts.append(f'[페이지 {i+1}]\n'+page.get_text('text')[:10000])
                if i in thumbs:
                    r=page.rect
                    if min(r.width,r.height)<=0:continue
                    pix=page.get_pixmap(matrix=fitz.Matrix(min(2,1120/max(r.width,r.height)),min(2,1120/max(r.width,r.height))),alpha=False)
                    with Image.open(io.BytesIO(pix.tobytes('png'))) as im:self.capture(im,s,f'{s["name"]} · {i+1}페이지','page',i+1)
                for f in page.get_fonts(full=True):add_font(self.result,s,f[3],'Bold' if re.search('bold',f[3],re.I) else '',f'PDF {i+1}페이지 폰트 리소스 (실사용 확인 필요)')
                if i in pages[:4]:
                    r=page.rect
                    self.result['dimensions'].append(dict(id=f'{s["id"]}-page{i+1}',sourceId=s['id'],label=f'{s["name"]} · {i+1}페이지',width=round(r.width*25.4/72,2),height=round(r.height*25.4/72,2),unit='mm',basis='PDF 페이지 크기 · 재단선/실물 크기와 다를 수 있음'))
            self.text(s,'\n\n'.join(texts),f'{len(pages)}/{doc.page_count}페이지 본문·폰트, {len(thumbs)}개 페이지 캡처 (동일 이미지 중복 제거)')
            if len(pages)<doc.page_count:s['status']='partial';self.warn(s,'전체 문서를 끝까지 읽지 않고 처음·중간·끝의 페이지를 균등 표본으로 읽었어요.')
            # A scan is NOT empty knowledge: rendered pages go to vision, without OCR calls.
            if not re.sub(r'\[페이지 \d+\]|\s','',s['text']):self.warn(s,'검색 가능한 본문이 없는 문서예요. AI 연결 시 캡처 이미지에서 보이는 내용을 확인해요.')
            embedded=0;seen=set()
            for i in thumbs:
                if embedded>=3:break
                for f in doc[i].get_images(full=True):
                    if embedded>=3:break
                    if f[0] in seen or min(f[2],f[3])<150:continue
                    seen.add(f[0])
                    try:
                        raw=doc.extract_image(f[0])['image']
                        with Image.open(io.BytesIO(raw)) as im:self.capture(im,s,f'{s["name"]} · {i+1}페이지 포함 이미지','embedded',i+1)
                        embedded+=1
                    except Exception:self.warn(s,'일부 포함 이미지의 추출을 생략했어요.')

    def raster(self,data,s):
        if s['ext'] in ('heic','heif'):
            try:from pillow_heif import register_heif_opener;register_heif_opener()
            except ImportError:raise ValueError('HEIC 해석 모듈이 없어요. 선택 설치 또는 JPG 변환이 필요해요.')
        with warnings.catch_warnings():
            warnings.simplefilter('error',Image.DecompressionBombWarning)
            with Image.open(io.BytesIO(data)) as im:
                if im.width*im.height>40_000_000:raise LimitError('이미지 해상도 한도를 초과했어요.')
                # PSD: use merged composite, never claim the parser rendered all layers.
                is_psd=im.format=='PSD'
                total=1 if is_psd else max(1,min(int(getattr(im,'n_frames',1)),400));indices=sample_indices(total,6);s['status']='read'
                for i in indices:
                    if not is_psd:im.seek(i)
                    frame=ImageOps.exif_transpose(im.copy())
                    self.capture(frame,s,s['name']+(f' · {i+1}프레임' if total>1 else ''),'frame' if total>1 else 'image',i+1 if total>1 else None)
                size=ImageOps.exif_transpose(im).size
                self.result['dimensions'].append(dict(id=s['id']+'-size',sourceId=s['id'],label=s['name'],width=size[0],height=size[1],unit='px',basis='이미지 디코더로 확인한 픽셀 크기'))
                s['coverage']=f'이미지 실제 픽셀·규격, {len(indices)}/{total}프레임 캡처'
                if total>1:s['status']='partial';self.warn(s,'움직임 전체를 재생·평가하지 않고 프레임을 표본 추출했어요.')
                if s['ext'] in ('psd','psb'):s['coverage']='PSD 저장된 합성 이미지·픽셀 크기';s['status']='partial';self.warn(s,'레이어 전체의 편집 상태는 판단하지 않았어요.')
                software=str(im.info.get('Software',''))
                try:software=software or str(im.getexif().get(305,''))
                except Exception:pass
                if software:s['text']='파일 메타데이터의 작성 소프트웨어: '+software[:120]

    def svg(self,data,s):
        root=ET.fromstring(data)
        if root.tag.split('}')[-1]!='svg':raise ValueError('SVG 루트가 아니에요.')
        # Only static markup. Never load a URL, script or foreignObject into a browser.
        removed=0
        for parent in root.iter():
            for child in list(parent):
                if child.tag.split('}')[-1] in ('script','foreignObject','image','use','animate','animateTransform','set'):
                    parent.remove(child);removed+=1
            for key,value in list(parent.attrib.items()):
                if key.lower().startswith('on') or key.split('}')[-1] in ('href','src') or ('url(' in value and any(not t.strip().strip("\'\"").startswith('#') for t in re.findall(r'url\((.*?)\)',value))):
                    del parent.attrib[key];removed+=1
        raw=ET.tostring(root,encoding='utf-8')
        self.text(s,decode_text(raw),'SVG 정적 선언·규격·폰트·벡터 캡처')
        # MuPDF SVG renderer is local; embedded/external resources removed above.
        with fitz.open(stream=raw,filetype='svg') as doc:
            page=doc[0];factor=min(2,1120/max(page.rect.width,page.rect.height));pix=page.get_pixmap(matrix=fitz.Matrix(factor,factor),alpha=False)
            with Image.open(io.BytesIO(pix.tobytes('png'))) as im:self.capture(im,s,s['name']+' · 정적 SVG 캡처','vector')
        w=re.fullmatch(r'([\d.]+)(px|mm|cm)?',root.get('width',''));h=re.fullmatch(r'([\d.]+)(px|mm|cm)?',root.get('height',''))
        if w and h and (w[2] or 'px')==(h[2] or 'px'):
            self.result['dimensions'].append(dict(id=s['id']+'-size',sourceId=s['id'],label=s['name'],width=float(w[1]),height=float(h[1]),unit=w[2] or 'px',basis='SVG width/height 선언'))
        if removed:s['status']='partial';self.warn(s,'안전을 위해 스크립트·외부/포함 이미지·동적 요소를 제외한 정적 SVG만 캡처했어요.')

    def embedded(self,z,s,prefixes,limit=6):
        infos=[i for i in z.infolist() if is_safe_member(i) and any(i.filename.startswith(p) for p in prefixes) and PurePosixPath(i.filename).suffix.lower().lstrip('.') in IMAGE_EXT-{'psb','psd'}]
        infos.sort(key=lambda i:(not any(k in i.filename.lower() for k in ('preview','thumbnail','mergedimage')),i.filename))
        for info in infos[:limit]:
            try:
                with Image.open(io.BytesIO(z.read(info))) as im:self.capture(im,s,s['name']+' · '+info.filename,'embedded',member=info.filename)
            except Exception:self.warn(s,'해석하지 못한 포함 이미지가 있어요.')
        return len(infos)

    def office(self,data,s,z=None):
        with safe_zip(data) as z:
            names=z.namelist();text=[];format=s['format']
            if format=='docx':
                parts=['word/document.xml']+[n for n in names if re.fullmatch(r'word/(header|footer)\d+\.xml',n)]
                for name in parts:
                    if name not in names:continue
                    root=ET.fromstring(z.read(name))
                    text += [''.join(n.text or '' for n in p.iter() if n.tag.endswith('}t')) for p in root.iter() if p.tag.endswith('}p')]
                font_parts=[n for n in ['word/document.xml','word/styles.xml'] if n in names]
                for name in font_parts:
                    root=ET.fromstring(z.read(name))
                    for node in root.iter():
                        if node.tag.endswith('}rFonts'):
                            for key,v in node.attrib.items():
                                if key.split('}')[-1] in ('ascii','hAnsi','eastAsia'):add_font(self.result,s,v,basis='DOCX 폰트 선언 · 미사용 스타일 포함 가능')
                self.embedded(z,s,('word/media/','docProps/thumbnail'))
            elif format=='pptx':
                slides=sorted([n for n in names if re.fullmatch('ppt/slides/slide[0-9]+.xml',n)],key=lambda x:int(re.search(r'(\d+)\.xml',x)[1]))
                chosen=sample_indices(len(slides),40)
                for i in chosen:
                    root=ET.fromstring(z.read(slides[i]));text.append(f'[슬라이드 {i+1}]\n'+'\n'.join(n.text or '' for n in root.iter() if n.tag.endswith('}t')))
                    for node in root.iter():
                        if node.tag.endswith(('}latin','}ea','}cs')) and node.get('typeface'):add_font(self.result,s,node.get('typeface'),basis='PPTX 슬라이드 폰트 선언')
                s['pagesTotal']=len(slides);s['pagesRead']=len(chosen);s['pageNumbersRead']=[i+1 for i in chosen]
                self.embedded(z,s,('ppt/media/','docProps/thumbnail'))
            elif format in ('xlsx','xlsm'):
                return self.sheet(data,s)
            elif format=='hwpx':
                for name in sorted(n for n in names if re.fullmatch(r'Contents/section\d+\.xml',n))[:40]:
                    root=ET.fromstring(z.read(name));text.append('\n'.join(''.join(p.itertext()) for p in root.iter() if p.tag.split('}')[-1]=='t'))
                self.embedded(z,s,('BinData/','Preview/'))
            elif format in ('odt','odp','ods'):
                root=ET.fromstring(z.read('content.xml'))
                text=[''.join(p.itertext()) for p in root.iter() if p.tag.split('}')[-1] in ('p','h')]
                self.embedded(z,s,('Pictures/','Thumbnails/'))
            elif format=='sketch':
                for name in [n for n in names if n=='meta.json' or n=='document.json' or n.startswith('pages/') and n.endswith('.json')][:30]:
                    obj=json.loads(z.read(name));stack=[obj];visited=0
                    while stack and visited<50000:
                        item=stack.pop();visited+=1
                        if isinstance(item,dict):
                            for k,v in item.items():
                                if k in ('string','name') and isinstance(v,str):text.append(v[:500])
                                if k=='NSFontNameAttribute' and isinstance(v,str):add_font(self.result,s,v,basis='Sketch 폰트 선언')
                                if isinstance(v,(list,dict)):stack.append(v)
                        elif isinstance(item,list):stack.extend(item[:2000])
                self.embedded(z,s,('previews/','images/'),8)
            else:
                self.embedded(z,s,('preview','Preview','mergedimage','Thumbnails/','Data/','QuickLook/'),8)
                self.text(s,'','저장된 미리보기·포함 이미지만 확인');s['status']='partial'
                self.warn(s,'독자적인 내부 편집 데이터를 완전히 해석하지 않았어요.');return
            self.text(s,'\n'.join(text),format.upper()+' 본문·선언 정보·포함 이미지 확인')
            if format in ('pptx','docx','odt','odp'):
                rendered=self.convert_office(data,s,z)
                if not rendered:
                    self.warn(s,'원본 페이지/슬라이드 캡처는 LibreOffice 설치와 ENABLE_OFFICE_RENDER=1 설정이 필요해요. 현재 본문과 포함 이미지를 확인했어요.')
                    s['status']='partial'
            if format=='hwpx':s['status']='partial';self.warn(s,'HWPX는 문단과 포함 이미지를 확인했으며 페이지 배치를 렌더링하지 않았어요.')

    def convert_office(self,data,s,z):
        tool=find_office()
        if not tool or not office_enabled() or self.conversions>=3:return False
        # Do not convert containers with active objects or externally linked resources.
        for info in z.infolist():
            low=info.filename.lower()
            if any(k in low for k in ('vbaproject','embeddings/','activex/','scripts/')):
                self.warn(s,'매크로/활성 개체가 포함되어 외부 변환기를 실행하지 않았어요.');return False
            if low.endswith(('.rels','.xml')) and info.file_size<MAX_MEMBER:
                raw=z.read(info)
                if re.search(br'TargetMode\s*=\s*[\'"]External|(?:xlink:href|href)\s*=\s*[\'"](?:https?:|file:|ftp:)',raw,re.I):
                    self.warn(s,'외부 참조가 있어 페이지 변환을 건너뛰었어요. 본문·포함 이미지는 유지해요.');return False
        self.conversions+=1
        with tempfile.TemporaryDirectory(prefix='morogaro-office-') as folder:
            root=Path(folder);src=root/('input.'+s['format']);src.write_bytes(data);profile=root/'profile';profile.mkdir()
            (profile/'user').mkdir()
            (profile/'user'/'registrymodifications.xcu').write_text('''<?xml version="1.0"?><oor:items xmlns:oor="http://openoffice.org/2001/registry"><item oor:path="/org.openoffice.Office.Common/Security/Scripting"><prop oor:name="MacroSecurityLevel" oor:op="fuse"><value>3</value></prop></item><item oor:path="/org.openoffice.Office.Common/Security/Scripting"><prop oor:name="DisableMacrosExecution" oor:op="fuse"><value>true</value></prop></item></oor:items>''')
            try:
                subprocess.run([tool,'-env:UserInstallation='+profile.as_uri(),'--headless','--nologo','--nodefault','--norestore','--convert-to','pdf','--outdir',str(root),str(src)],stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL,timeout=25,env=scrub_env(root),check=True)
                pdf=root/'input.pdf'
                if not pdf.is_file() or pdf.stat().st_size>MAX_MEMBER:raise ValueError('conversion')
                saved_text=s['text'];self.pdf(pdf.read_bytes(),s);s['text']=bounded_text(saved_text+'\n'+s['text'])
                s['coverage']+=' · LibreOffice 변환 캡처 (폰트 환경에 따라 원본과 차이 가능)'
                return True
            except Exception:self.warn(s,'페이지 변환에 실패해 본문·포함 이미지 결과로 이어갔어요.');return False

    def sheet(self,data,s):
        import openpyxl
        # Read-only cached values. No formula evaluation, external links or VBA.
        wb=openpyxl.load_workbook(io.BytesIO(data),read_only=True,data_only=True,keep_links=False)
        try:
            alltext=[]
            for sheet in wb.worksheets[:8]:
                rows=[]
                for cells in sheet.iter_rows(min_row=1,max_row=min(sheet.max_row or 200,200),max_col=min(sheet.max_column or 20,20),values_only=True):
                    vals=[str(c)[:400] if c is not None else '' for c in cells]
                    if any(vals):rows.append(' | '.join(vals))
                alltext.append('[시트 '+sheet.title+']\n'+'\n'.join(rows))
                if rows:self.capture(text_card(sheet.title,rows),s,s['name']+' · '+sheet.title+' · 표 데이터 미리보기','table',member=sheet.title,sample_colors=False)
            self.text(s,'\n\n'.join(alltext),'앞 8시트의 시트당 200행×20열 캐시 값 · 표 데이터 미리보기')
            s['status']='partial';self.warn(s,'수식은 실행하지 않아요. 저장된 계산값만 읽으며 차트·원본 레이아웃·색은 분석하지 않았어요.')
        finally:wb.close()

    def hwp(self,data,s):
        try:import olefile
        except ImportError:raise ValueError('HWP 해석을 위한 olefile 모듈이 없어요.')
        with olefile.OleFileIO(io.BytesIO(data)) as ole:
            header=ole.openstream('FileHeader').read(256)
            if not header.startswith(b'HWP Document File'):raise ValueError('HWP 문서가 아니에요.')
            flags=struct.unpack_from('<I',header,36)[0]
            if flags&(2|4):raise ValueError('암호·배포용 보호 문서는 읽지 않았어요.')
            texts=[]
            streams=[p for p in ole.listdir() if p[0]=='BodyText' and p[-1].startswith('Section')][:40]
            for path in streams:
                data=ole.openstream(path).read(12*1024*1024)
                if flags&1:
                    obj=zlib.decompressobj(-15);data=obj.decompress(data,12*1024*1024)
                    if obj.unconsumed_tail:raise LimitError('HWP 본문 압축 한도 초과')
                offset=0
                while offset+4<=len(data):
                    record=struct.unpack_from('<I',data,offset)[0];offset+=4;tag=record&0x3ff;size=record>>20
                    if size==0xfff:
                        if offset+4>len(data):break
                        size=struct.unpack_from('<I',data,offset)[0];offset+=4
                    if offset+size>len(data):break
                    chunk=data[offset:offset+size];offset+=size
                    if tag==67:
                        text=chunk.decode('utf-16le',errors='replace');texts.append(re.sub(r'[\x00-\x08\x0b-\x1f]',' ',text))
            if not texts and ole.exists('PrvText'):texts=[ole.openstream('PrvText').read(128*1024).decode('utf-16le',errors='replace')]
            if ole.exists('PrvImage'):
                try:
                    with Image.open(io.BytesIO(ole.openstream('PrvImage').read(MAX_MEMBER))) as im:self.capture(im,s,s['name']+' · 저장된 미리보기','embedded')
                except Exception:pass
            self.text(s,'\n'.join(texts),'HWP5 문단 텍스트·저장된 미리보기 (있을 때)');s['status']='partial';self.warn(s,'HWP 페이지 레이아웃·전체 오브젝트는 렌더링하지 않았어요.')

    def media(self,data,s):
        ffmpeg,ffprobe=shutil.which('ffmpeg'),shutil.which('ffprobe')
        magic=data[:32]
        recognized=(magic[4:8] in (b'ftyp',b'moov',b'mdat',b'wide') or magic.startswith((b'\x1aE\xdf\xa3',b'ID3',b'fLaC',b'OggS',b'\x00\x00\x01')) or (magic.startswith(b'RIFF') and magic[8:12] in (b'AVI ',b'WAVE')) or (len(magic)>1 and magic[0]==255 and magic[1]&224==224))
        if not recognized:raise ValueError('허용된 미디어 헤더가 아니에요. 재생목록·외부 참조 파일은 해석하지 않아요.')
        if not ffmpeg or not ffprobe:raise ValueError('영상·음성 해석에는 FFmpeg 설치가 필요해요.')
        with tempfile.TemporaryDirectory(prefix='morogaro-media-') as folder:
            root=Path(folder);path=root/('input.'+s['ext']);path.write_bytes(data);env=scrub_env(root)
            probe=subprocess.run([ffprobe,'-v','error','-protocol_whitelist','file,pipe','-show_entries','format=duration:stream=codec_type,width,height,codec_name','-of','json',str(path)],capture_output=True,timeout=10,check=True,env=env)
            if len(probe.stdout)>100000:raise LimitError('미디어 메타데이터 한도 초과')
            meta=json.loads(probe.stdout);duration=float(meta.get('format',{}).get('duration',0));duration=duration if math.isfinite(duration) and duration>=0 else 0
            streams=meta.get('streams',[]);video=next((v for v in streams if v.get('codec_type')=='video'),None);audio=any(v.get('codec_type')=='audio' for v in streams)
            s['duration']=round(duration,2);s['status']='partial';s['coverage']='미디어 메타데이터'
            if video:
                if int(video.get('width',0))*int(video.get('height',0))>40_000_000:raise LimitError('영상 해상도 한도 초과')
                stamps=[round(duration*p,2) for p in (.03,.18,.36,.55,.74,.93)] if duration>1 else [0]
                for i,ts in enumerate(stamps):
                    output=root/f'frame-{i}.jpg'
                    try:
                        subprocess.run([ffmpeg,'-v','error','-nostdin','-threads','1','-protocol_whitelist','file,pipe','-ss',str(ts),'-i',str(path),'-an','-sn','-frames:v','1','-vf',"scale=960:960:force_original_aspect_ratio=decrease",'-q:v','3','-y',str(output)],stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL,timeout=10,check=True,env=env)
                        with Image.open(output) as im:self.capture(im,s,f'{s["name"]} · {int(ts//60):02d}:{int(ts%60):02d}','video',seconds=ts)
                    except Exception:self.warn(s,'일부 시점의 영상 프레임을 캡처하지 못했어요.')
                s['coverage']=f'전체 {duration:.1f}초 중 {len(stamps)}시점 정지 프레임·규격'
                s['frameTimes']=stamps
                self.result['dimensions'].append(dict(id=s['id']+'-video',sourceId=s['id'],label=s['name'],width=video.get('width'),height=video.get('height'),unit='px',basis='영상 스트림 해상도'))
            if audio and self.audio_count<MAX_AUDIO_FILES:
                audiofile=root/'sample.wav'
                subprocess.run([ffmpeg,'-v','error','-nostdin','-threads','1','-protocol_whitelist','file,pipe','-i',str(path),'-t',str(MAX_AUDIO_SECONDS),'-vn','-sn','-ac','1','-ar','16000','-c:a','pcm_s16le','-y',str(audiofile)],stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL,timeout=15,check=True,env=env)
                if audiofile.is_file() and audiofile.stat().st_size<5*1024*1024:
                    s['_audio']=base64.b64encode(audiofile.read_bytes()).decode();s['audioSeconds']=min(duration or MAX_AUDIO_SECONDS,MAX_AUDIO_SECONDS);self.audio_count+=1
                    self.warn(s,'음성은 AI 이용에 동의한 경우에만 앞 120초 이내를 전사해요. 음악·전체 영상의 이해를 보장하지 않아요.')
            elif audio:self.warn(s,'음성 추출은 요청당 최대 3파일까지만 진행해요.')

    def archive(self,data,s,depth):
        with safe_zip(data) as z:
            entries=[i for i in z.infolist() if not i.is_dir()];handled=0
            for i in entries:
                if len(self.result['sources'])>=MAX_SOURCES or handled>=40:
                    self.warn(s,'압축 내부 파일 개수 한도에 도달해 나머지는 읽지 않았어요.');break
                if not is_safe_member(i):self.warn(s,'암호·경로 이탈·심볼릭 링크 항목을 제외했어요.');continue
                path=i.filename.replace('\\','/');parts=PurePosixPath(path).parts
                if any(p.startswith('.') or p in ('__MACOSX','node_modules','.git','vendor','__pycache__') for p in parts):continue
                if PurePosixPath(path).suffix.lower().lstrip('.') in NEVER_AUTO:continue
                if self.expanded+i.file_size>MAX_EXPANDED_BYTES:raise LimitError('전체 압축 해제 용량 한도 초과')
                self.expanded+=i.file_size;raw=z.read(i);handled+=1
                self.file({'name':path,'data':raw},s['id']+f'-a{handled}',parent=s['id'],depth=depth+1)
            s['status']='partial';s['coverage']=f'압축 내부 {handled}/{len(entries)}항목 확인 · 숨김 파일·실행파일 제외'
            s['text']='압축 내부 파일 목록:\n'+'\n'.join(i.filename for i in entries[:80])
            s['childrenCount']=handled

    def file(self,f,id,parent=None,depth=0):
        name=clean_name(f['name']);ext=PurePosixPath(name).suffix.lower().lstrip('.');data=f['data']
        s=dict(id=id,name=name,ext=ext,size=len(data),kind=KIND.get(ext,'코드' if ext in TEXT_EXT else '기타 파일'),kindBasis='확장자·컨테이너 기준 · 직접 수정 가능',status='metadata',coverage='파일명·형식·용량만 확인',text='',warnings=[],parentId=parent,format=ext,sha256=hashlib.sha256(data).hexdigest())
        self.result['sources'].append(s);self.progress(dict(phase='extract',file=name,processed=len(self.result['sources']),message='파일 내용을 읽고 썸네일 후보를 캡처하고 있어요.'))
        if ext in TOOL_MAP:
            self.result['values']['tools']=uniq(self.result['values']['tools']+[TOOL_MAP[ext]]);evidence(self.result,'tools',s,ext+' 확장자에 연결된 툴 후보','extension','파일 형식')
        try:
            if len(data)>MAX_FILE:raise LimitError('파일당 50MB 분석 한도를 초과했어요.')
            if PurePosixPath(name).name in ('.env','.npmrc','.pypirc','credentials') or (ext in NEVER_AUTO and ext!='key'):raise ValueError('실행파일·인증정보·폰트 바이너리는 분석 대상에서 제외했어요.')
            if b'%PDF-' in data[:1024]:s['format']='pdf';self.pdf(data,s)
            elif data.startswith(b'PK'):
                with safe_zip(data) as z:
                    names=z.namelist()
                    if 'word/document.xml' in names:s['format']='docx'
                    elif 'ppt/presentation.xml' in names or any(n.startswith('ppt/slides/slide') for n in names):s['format']='pptx'
                    elif 'xl/workbook.xml' in names:s['format']='xlsx'
                    elif any(re.fullmatch(r'Contents/section\d+\.xml',n) for n in names):s['format']='hwpx'
                    elif 'content.xml' in names:s['format']=ext if ext in ('odt','ods','odp') else 'odt'
                    elif 'document.json' in names and 'meta.json' in names:s['format']='sketch'
                    elif ext in ('kra','ora','pages','key','numbers'):s['format']=ext
                    else:s['format']='zip'
                if s['format']=='zip':
                    if depth>=MAX_ARCHIVE_DEPTH:raise LimitError('중첩 ZIP은 최대 2단계까지 확인해요.')
                    self.archive(data,s,depth)
                else:self.office(data,s)
            elif ext=='svg' or data[:300].lstrip().startswith(b'<svg'):s['format']='svg';self.svg(data,s)
            elif data.startswith(b'\xd0\xcf\x11\xe0') and ext=='hwp':self.hwp(data,s)
            elif ext in VIDEO_EXT|AUDIO_EXT:self.media(data,s)
            elif ext in IMAGE_EXT or data.startswith((b'\x89PNG',b'\xff\xd8',b'GIF8',b'8BPS',b'II*\x00',b'MM\x00*')):self.raster(data,s)
            elif ext in TEXT_EXT or b'\0' not in data[:1000] and not ext and data:
                text=decode_text(data)
                if ext in ('html','htm'):
                    # Text/inline declarations only: no browser, remote CSS or JS.
                    text=re.sub(r'<(script|style)\b[^>]*>[\s\S]*?</\1>','',text,flags=re.I)
                self.text(s,text,'본문·선언 데이터 (코드 실행 없음)')
                if len(data)>256*1024:self.warn(s,'텍스트 입력은 앞 256KB를 읽었으며 이후 내용은 제외했어요.')
            elif ext=='fig':self.warn(s,'Figma 로컬 .fig의 비공개 바이너리를 해석하지 않았어요. 프레임을 PDF/PNG/SVG로 함께 내보내 주세요.')
            else:self.warn(s,'이 내부 형식은 아직 해석하지 못해요. PDF·이미지·문서로 내보낸 자료를 함께 올려주세요.')
        except Exception as exc:
            s['status']='partial' if s.get('text') or self.previews_per_source[id] else 'metadata'
            s['errorType']=type(exc).__name__
            if isinstance(exc,(ValueError,LimitError)) and str(exc) and len(str(exc))<200:self.warn(s,str(exc))
            else:self.warn(s,'손상·암호·형식·처리 한도 때문에 내용을 모두 읽지 못했어요.')
        if s['format']!=ext:self.warn(s,'확장자와 실제 내부 형식이 달라 '+s['format'].upper()+' 방식으로 확인했어요.')
        if ext=='ai' and s['format']=='pdf':
            s['status']='partial';self.warn(s,'Illustrator의 PDF 호환 부분만 읽었어요. 편집 레이어·전체 AI 내부 정보는 해석하지 않았어요.')
        s['thumbnailCount']=self.previews_per_source[id]
        if s['status']=='read' and any(any(k in w for k in ('한도','생략','제외','못했')) for w in s['warnings']):s['status']='partial'
        return s

    def finish(self,files):
        for i,f in enumerate(files[:MAX_FILES]):self.file(f,'f'+str(i))
        if len(files)>MAX_FILES:self.result['warnings'].append('상위 파일은 요청당 20개까지만 해석해요.')
        # Candidates are ranked, not cropped/generated by AI. Selection remains user editable.
        self.result['thumbnails'].sort(key=lambda p:p['quality'],reverse=True)
        self.result['recommendedThumbnails']=[p['id'] for p in self.result['thumbnails'][:6]]
        self.result['coverage']=dict(uploadedFiles=len(files),inspectedSources=len(self.result['sources']),thumbnailCandidates=len(self.result['thumbnails']),fullyRead=sum(s['status']=='read' for s in self.result['sources']),partiallyRead=sum(s['status']=='partial' for s in self.result['sources']),unread=sum(s['status']=='metadata' for s in self.result['sources']))
        self.result['pipelineVersion']=PIPELINE_VERSION
        return self.result
